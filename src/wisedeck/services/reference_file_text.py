"""Extract plain text from uploaded reference files (light + optional MarkItDown deep path)."""

from __future__ import annotations

import io
import logging
import re
from pathlib import Path
from typing import Literal

logger = logging.getLogger(__name__)

MAX_STORE_CHARS = 120_000

ParseMode = Literal["light", "deep"]


def _extract_docx_paragraphs_and_tables(content: bytes) -> str:
    """Paragraphs plus tables as tab-separated blocks (markdown-friendly)."""
    try:
        import docx

        doc = docx.Document(io.BytesIO(content))
        parts: list[str] = []
        for p in doc.paragraphs:
            t = (p.text or "").strip()
            if t:
                parts.append(t)
        for table in doc.tables:
            rows_out: list[str] = []
            for row in table.rows:
                cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    rows_out.append("\t".join(cells))
            if rows_out:
                parts.append("### TABLE\n" + "\n".join(rows_out))
        return "\n".join(parts)
    except Exception as e:
        logger.warning("DOCX structured extract failed: %s", e)
        return ""


def extract_text_light(content: bytes, filename: str) -> str:
    """Fast path: PyPDF2 / docx / text decode."""
    suffix = Path(filename or "").suffix.lower()
    text = ""
    try:
        if suffix == ".pdf":
            try:
                import PyPDF2

                reader = PyPDF2.PdfReader(io.BytesIO(content))
                parts = []
                for page in reader.pages:
                    t = page.extract_text() or ""
                    if t.strip():
                        parts.append(t.strip())
                text = "\n".join(parts)
            except Exception as e:
                logger.warning("PDF text extract failed for %s: %s", filename, e)
                text = ""
        elif suffix == ".docx":
            text = _extract_docx_paragraphs_and_tables(content)
        elif suffix in (".md", ".txt", ".html", ".htm", ".csv", ".json"):
            for enc in ("utf-8", "gbk", "latin-1"):
                try:
                    text = content.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                text = content.decode("utf-8", errors="replace")
        else:
            text = content.decode("utf-8", errors="replace")
            if len(text.strip()) < 40 and suffix in (".pptx", ".xlsx", ".bin"):
                text = ""
                logger.info("Reference file %s looks binary; no text extracted", filename)
    except Exception as e:
        logger.warning("Reference extract error for %s: %s", filename, e)
        text = ""

    text = (text or "").strip()
    if len(text) > MAX_STORE_CHARS:
        text = text[:MAX_STORE_CHARS] + "\n…(已截断存储)"
    return text


def extract_text_markitdown_path(storage_path: str) -> str:
    """Deep path using MarkItDown on an on-disk file (supports many Office/PDF types)."""
    try:
        from markitdown import MarkItDown

        result = MarkItDown().convert(storage_path)
        raw = getattr(result, "text_content", None) or getattr(result, "markdown", None) or ""
        text = str(raw).strip()
        if len(text) > MAX_STORE_CHARS:
            text = text[:MAX_STORE_CHARS] + "\n…(已截断存储)"
        return text
    except Exception as e:
        logger.warning("MarkItDown convert failed for %s: %s", storage_path, e)
        return ""


_TOKEN_RE = re.compile(r"[\w\u4e00-\u9fff]+", re.UNICODE)


def _tokenize_for_score(text: str) -> set[str]:
    return {t.lower() for t in _TOKEN_RE.findall(text or "") if len(t) > 1}


def extract_text_from_upload(
    content: bytes,
    filename: str,
    *,
    storage_path: str | None = None,
    parse_mode: ParseMode = "light",
) -> tuple[str, str]:
    """
    Returns (text, mode_used) where mode_used is 'light' or 'deep'.
    When parse_mode is deep, tries MarkItDown on storage_path then picks richer result.
    """
    light = extract_text_light(content, filename)
    if parse_mode != "deep" or not storage_path:
        return light, "light"
    deep = extract_text_markitdown_path(storage_path)
    if len(deep.strip()) > len(light.strip()):
        return deep, "deep"
    if len(light.strip()) < 80 and len(deep.strip()) >= 80:
        return deep, "deep"
    if light.strip():
        return light, "light"
    if deep.strip():
        return deep, "deep"
    return light, "light"
