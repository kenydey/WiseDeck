"""
PPTX physical structure extractor (python-pptx).

Goal: extract "what boxes exist and where" (layout + shapes + geometry + text-frame flags),
without caring about specific text content.
"""

from __future__ import annotations

import json
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

_SCHEMA_VERSION = 1
_MAX_SLIDES = 80
_MAX_SHAPES_PER_SLIDE = 80
_MAX_JSON_BYTES = 380_000


def _round4(x: float) -> float:
    return float(f"{x:.4f}")


def _safe_int(v: Any) -> int:
    try:
        return int(v or 0)
    except Exception:
        return 0


def _safe_bool(v: Any) -> Optional[bool]:
    if v is None:
        return None
    try:
        return bool(v)
    except Exception:
        return None


def _text_frame_flags(shape: Any) -> Optional[Dict[str, Any]]:
    if not getattr(shape, "has_text_frame", False):
        return None
    tf = getattr(shape, "text_frame", None)
    if tf is None:
        return None
    out: Dict[str, Any] = {}
    try:
        out["word_wrap"] = _safe_bool(getattr(tf, "word_wrap", None))
    except Exception:
        pass
    # python-pptx names autofit as `auto_size` (MSO_AUTO_SIZE enum)
    try:
        auto = getattr(tf, "auto_size", None)
        if auto is not None:
            out["auto_size"] = str(auto).split(".")[-1]
    except Exception:
        pass
    return out or None


def _placeholder_type(shape: Any) -> Optional[str]:
    try:
        if not bool(getattr(shape, "is_placeholder", False)):
            return None
        pf = getattr(shape, "placeholder_format", None)
        if pf is None:
            return None
        t = getattr(pf, "type", None)
        if t is None:
            return None
        return str(t).split(".")[-1]
    except Exception:
        return None


def _shape_kind(shape: Any) -> str:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        st = getattr(shape, "shape_type", None)
        if st == MSO_SHAPE_TYPE.CHART:
            return "chart"
        if st == MSO_SHAPE_TYPE.TABLE:
            return "table"
        if st == MSO_SHAPE_TYPE.PICTURE:
            return "picture"
        if getattr(shape, "has_text_frame", False):
            # includes textbox / autoshape with text
            return "text"
        return "other"
    except Exception:
        return "other"


def _normalized_bbox(shape: Any, *, slide_w: int, slide_h: int) -> Optional[List[float]]:
    if slide_w <= 0 or slide_h <= 0:
        return None
    left = _safe_int(getattr(shape, "left", 0))
    top = _safe_int(getattr(shape, "top", 0))
    width = _safe_int(getattr(shape, "width", 0))
    height = _safe_int(getattr(shape, "height", 0))
    try:
        return [
            _round4(left / slide_w),
            _round4(top / slide_h),
            _round4(width / slide_w),
            _round4(height / slide_h),
        ]
    except Exception:
        return None


def _static_signature(
    *,
    kind: str,
    bbox_norm: Optional[List[float]],
    layout_name: str,
) -> Optional[str]:
    """
    A coarse signature for cross-slide repetition heuristic.
    - only for non-placeholder shapes
    - quantize bbox to 4% grid
    """
    if not bbox_norm or len(bbox_norm) < 4:
        return None
    l, t, w, h = bbox_norm[:4]
    q = lambda v: round(float(v) / 0.04) * 0.04  # noqa: E731
    try:
        return f"{layout_name}|{kind}|x{q(l):.2f}|y{q(t):.2f}|w{q(w):.2f}|h{q(h):.2f}"
    except Exception:
        return None


def extract_pptx_physical_structure(pptx_bytes: bytes) -> Dict[str, Any]:
    """
    Return a JSON-serializable dict:
      schema_version, slide_width_emu, slide_height_emu,
      slides: [{index, layout_name, shapes:[...]}]

    shapes[] contains:
      shape_id, kind, is_placeholder, placeholder_type, bbox_emu, bbox, text_frame(optional),
      is_static(optional), static_role(optional)
    """
    try:
        from pptx import Presentation
    except Exception as e:
        return {"schema_version": _SCHEMA_VERSION, "error": f"python-pptx unavailable: {e}", "slides": []}

    if not pptx_bytes or not isinstance(pptx_bytes, (bytes, bytearray)):
        return {"schema_version": _SCHEMA_VERSION, "error": "empty_bytes", "slides": []}

    try:
        prs = Presentation(BytesIO(pptx_bytes))
    except Exception as e:
        return {"schema_version": _SCHEMA_VERSION, "error": str(e)[:200], "slides": []}

    sw = _safe_int(getattr(prs, "slide_width", 0))
    sh = _safe_int(getattr(prs, "slide_height", 0))

    slides_list: List[Any] = []
    try:
        slides_list = list(prs.slides)[:_MAX_SLIDES]
    except Exception:
        slides_list = []

    # First pass: gather signatures for potential static elements.
    sig_counts: Dict[str, int] = {}
    sig_samples: Dict[str, Dict[str, Any]] = {}

    slides_out: List[Dict[str, Any]] = []
    for si, slide in enumerate(slides_list, start=1):
        layout_name = ""
        try:
            lo = getattr(slide, "slide_layout", None)
            layout_name = str(getattr(lo, "name", "") or "")
        except Exception:
            layout_name = ""

        shapes_payload: List[Dict[str, Any]] = []
        try:
            shapes = list(getattr(slide, "shapes", []) or [])[:_MAX_SHAPES_PER_SLIDE]
        except Exception:
            shapes = []

        for shape in shapes:
            shape_id = _safe_int(getattr(shape, "shape_id", 0))
            is_ph = False
            try:
                is_ph = bool(getattr(shape, "is_placeholder", False))
            except Exception:
                is_ph = False

            ph_type = _placeholder_type(shape)
            kind = _shape_kind(shape)
            bbox_norm = _normalized_bbox(shape, slide_w=sw, slide_h=sh)
            bbox_emu = [
                _safe_int(getattr(shape, "left", 0)),
                _safe_int(getattr(shape, "top", 0)),
                _safe_int(getattr(shape, "width", 0)),
                _safe_int(getattr(shape, "height", 0)),
            ]

            entry: Dict[str, Any] = {
                "shape_id": shape_id,
                "kind": kind,
                "is_placeholder": bool(is_ph),
                "placeholder_type": ph_type,
                "bbox_emu": bbox_emu,
                "bbox": bbox_norm,
            }
            tf = _text_frame_flags(shape)
            if tf:
                entry["text_frame"] = tf

            # Static heuristic signature (non-placeholder)
            if not is_ph:
                sig = _static_signature(kind=kind, bbox_norm=bbox_norm, layout_name=layout_name or "")
                if sig:
                    sig_counts[sig] = sig_counts.get(sig, 0) + 1
                    if sig not in sig_samples:
                        sig_samples[sig] = {
                            "kind": kind,
                            "bbox": bbox_norm,
                            "layout_name": layout_name,
                        }
                    entry["_static_sig"] = sig

            shapes_payload.append(entry)

        slides_out.append({"index": si, "layout_name": layout_name, "shapes": shapes_payload})

    # Second pass: mark static if repeated across many slides (or matches header/footer zones).
    slide_n = len(slides_out) or 0
    if slide_n >= 2:
        threshold = max(2, int(slide_n * 0.6 + 0.999))
        static_sigs = {s for s, c in sig_counts.items() if c >= threshold}
    else:
        static_sigs = set()

    def _zone_role(bbox: Optional[List[float]]) -> Optional[str]:
        if not bbox or len(bbox) < 4:
            return None
        l, t, w, h = bbox[:4]
        # simple header/footer zones
        if t <= 0.08 and h <= 0.20:
            return "header"
        if (t + h) >= 0.90 and h <= 0.20:
            # small footer-ish element
            if w <= 0.30:
                return "page_number"
            return "footer"
        return None

    for slide in slides_out:
        for sh_entry in slide.get("shapes") or []:
            if not isinstance(sh_entry, dict):
                continue
            if sh_entry.get("is_placeholder"):
                sh_entry.pop("_static_sig", None)
                continue
            sig = sh_entry.pop("_static_sig", None)
            bbox = sh_entry.get("bbox")
            role = _zone_role(bbox if isinstance(bbox, list) else None)
            if role:
                sh_entry["static_role"] = role
            if sig and sig in static_sigs:
                sh_entry["is_static"] = True
            elif role in {"header", "footer", "page_number"}:
                # zone-based soft static mark (best-effort)
                sh_entry["is_static"] = True

    payload: Dict[str, Any] = {
        "schema_version": _SCHEMA_VERSION,
        "slide_width_emu": sw,
        "slide_height_emu": sh,
        "slide_count": len(slides_out),
        "slides": slides_out,
    }

    # Cap size by trimming slide tail.
    raw = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    if len(raw.encode("utf-8")) > _MAX_JSON_BYTES and slides_out:
        payload["slides"] = slides_out[: max(1, min(20, len(slides_out)))]
        payload["truncated"] = True
        payload["truncated_slide_count"] = len(slides_out) - len(payload["slides"])
    return payload


__all__ = ["extract_pptx_physical_structure"]

