"""Visual DNA extraction v1 (minimal) from PPTX via python-pptx.

Goal: provide a stable, small set of style signals for style extrapolation:
- paletteTop5
- fontPair (title/body)
- simple decoration hint (bottom bar)

This module intentionally avoids heavy dependencies and keeps heuristics simple.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple


def _hex_from_rgb(rgb: Any) -> Optional[str]:
    try:
        r, g, b = int(rgb[0]), int(rgb[1]), int(rgb[2])
        if not (0 <= r <= 255 and 0 <= g <= 255 and 0 <= b <= 255):
            return None
        return f"#{r:02X}{g:02X}{b:02X}"
    except Exception:
        return None


def _shape_fill_hex(shape: Any) -> Optional[str]:
    try:
        fill = getattr(shape, "fill", None)
        if fill is None:
            return None
        fore = getattr(fill, "fore_color", None)
        if fore is None:
            return None
        rgb = getattr(fore, "rgb", None)
        if rgb is None:
            return None
        return _hex_from_rgb(rgb)
    except Exception:
        return None


def _shape_line_hex(shape: Any) -> Optional[str]:
    try:
        line = getattr(shape, "line", None)
        if line is None:
            return None
        fc = getattr(line, "color", None)
        rgb = getattr(fc, "rgb", None) if fc is not None else None
        if rgb is None:
            return None
        return _hex_from_rgb(rgb)
    except Exception:
        return None


def _extract_text_font_info(shape: Any) -> Tuple[Optional[str], Optional[float]]:
    """Return (font_name, font_size_pt) from first run if available."""
    try:
        if not getattr(shape, "has_text_frame", False):
            return None, None
        tf = shape.text_frame
        for p in getattr(tf, "paragraphs", []) or []:
            for r in getattr(p, "runs", []) or []:
                f = getattr(r, "font", None)
                if f is None:
                    continue
                name = getattr(f, "name", None)
                size = getattr(f, "size", None)
                size_pt = float(size.pt) if size is not None and getattr(size, "pt", None) else None
                if name or size_pt:
                    return (str(name) if name else None), size_pt
    except Exception:
        return None, None
    return None, None


@dataclass
class VisualDNAV1:
    palette_top5: List[str]
    primary_color: Optional[str]
    font_title: Optional[str]
    font_body: Optional[str]
    font_scale_ratio: Optional[float]
    decoration_bottom_bar: Optional[Dict[str, Any]]
    warnings: List[str]

    def to_dict(self) -> Dict[str, Any]:
        return {
            "paletteTop5": list(self.palette_top5),
            "primaryColor": self.primary_color,
            "fontPair": {
                "title": self.font_title,
                "body": self.font_body,
                "title_body_ratio": self.font_scale_ratio,
            },
            "decoration": {"bottom_bar": self.decoration_bottom_bar},
            "warnings": list(self.warnings),
        }


def extract_visual_dna_v1_from_pptx_bytes(pptx_bytes: bytes) -> VisualDNAV1:
    warnings: List[str] = []
    colors: List[str] = []
    title_fonts: List[str] = []
    body_fonts: List[str] = []
    title_sizes: List[float] = []
    body_sizes: List[float] = []
    bottom_bar: Optional[Dict[str, Any]] = None

    try:
        from pptx import Presentation  # type: ignore
    except Exception as e:
        return VisualDNAV1(
            palette_top5=[],
            primary_color=None,
            font_title=None,
            font_body=None,
            font_scale_ratio=None,
            decoration_bottom_bar=None,
            warnings=[f"python-pptx unavailable: {str(e)[:200]}"],
        )

    prs = Presentation(BytesIO(pptx_bytes))
    sw = float(getattr(prs, "slide_width", 0) or 0)
    sh = float(getattr(prs, "slide_height", 0) or 0)
    if sw <= 0 or sh <= 0:
        warnings.append("invalid_slide_dimensions")
        sw, sh = 1.0, 1.0

    # Scan slides (cap to keep v1 fast)
    for si, slide in enumerate(list(prs.slides)[:12], start=1):
        try:
            for shape in list(getattr(slide, "shapes", []) or [])[:200]:
                c = _shape_fill_hex(shape) or _shape_line_hex(shape)
                if c:
                    colors.append(c)

                fn, fs = _extract_text_font_info(shape)
                if fn:
                    # Heuristic: large font tends to be title
                    if fs and fs >= 28:
                        title_fonts.append(fn)
                        title_sizes.append(fs)
                    else:
                        body_fonts.append(fn)
                        if fs:
                            body_sizes.append(fs)

                # Bottom bar heuristic: wide, short, near bottom, solid fill.
                if bottom_bar is None:
                    try:
                        w = float(getattr(shape, "width", 0) or 0)
                        h = float(getattr(shape, "height", 0) or 0)
                        l = float(getattr(shape, "left", 0) or 0)
                        t = float(getattr(shape, "top", 0) or 0)
                        fill_c = _shape_fill_hex(shape)
                        if fill_c and w / sw >= 0.6 and h / sh <= 0.08 and (t + h) / sh >= 0.9:
                            bottom_bar = {
                                "color": fill_c,
                                "bbox_pct": [l / sw, t / sh, w / sw, h / sh],
                                "slide_index": si,
                            }
                    except Exception:
                        pass
        except Exception as e:
            warnings.append(f"slide_scan_failed[{si}]: {str(e)[:120]}")

    # Palette: simple frequency-based Top-5 with dedupe
    freq: Dict[str, int] = {}
    for c in colors:
        freq[c] = freq.get(c, 0) + 1
    ordered = sorted(freq.items(), key=lambda kv: kv[1], reverse=True)
    palette: List[str] = []
    for c, _n in ordered:
        if c not in palette:
            palette.append(c)
        if len(palette) >= 5:
            break

    primary = palette[0] if palette else None

    def _most_common(xs: List[str]) -> Optional[str]:
        if not xs:
            return None
        f: Dict[str, int] = {}
        for x in xs:
            f[x] = f.get(x, 0) + 1
        return sorted(f.items(), key=lambda kv: kv[1], reverse=True)[0][0]

    title_font = _most_common(title_fonts) or _most_common(body_fonts)
    body_font = _most_common(body_fonts) or title_font

    ratio = None
    if title_sizes and body_sizes:
        try:
            ratio = round((sum(title_sizes) / len(title_sizes)) / (sum(body_sizes) / len(body_sizes)), 2)
        except Exception:
            ratio = None

    if len(palette) < 3:
        warnings.append("palette_low_confidence")
    if not title_font or not body_font:
        warnings.append("font_pair_low_confidence")

    return VisualDNAV1(
        palette_top5=palette,
        primary_color=primary,
        font_title=title_font,
        font_body=body_font,
        font_scale_ratio=ratio,
        decoration_bottom_bar=bottom_bar,
        warnings=warnings,
    )

