"""
Compact per-slide layout hints from PPTX (python-pptx) for import_summary / SVG injection.

Designed to stay small for DB JSON: normalized bbox, enum names only, no blob/text bodies.
"""

from __future__ import annotations

import json
from io import BytesIO
from typing import Any, Dict, List, Optional

_MAX_SLIDES = 48
_MAX_SHAPES_PER_SLIDE = 28
_MAX_JSON_BYTES = 120_000


def _ph_type_name(shape: object) -> Optional[str]:
    try:
        if not bool(getattr(shape, "is_placeholder", False)):
            return None
        pf = getattr(shape, "placeholder_format", None)
        if pf is None:
            return None
        t = getattr(pf, "type", None)
        if t is None:
            return None
        return str(t).split(".")[-1]
    except Exception:
        return None


def _shape_kind(shape: object) -> str:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        st = getattr(shape, "shape_type", None)
        if st == MSO_SHAPE_TYPE.CHART:
            return "chart"
        if st == MSO_SHAPE_TYPE.PICTURE:
            return "picture"
        if getattr(shape, "has_text_frame", False):
            return "text"
        return "other"
    except Exception:
        return "other"


def extract_pptx_layout_hints(pptx_bytes: bytes) -> Dict[str, Any]:
    """
    Return a JSON-serializable dict:
      schema_version, slide_width_emu, slide_height_emu,
      slides: [{index, shapes: [{bbox, placeholder_type, is_placeholder, shape_kind}]}]

    bbox is [left, top, width, height] each in 0..1 relative to slide size.
    """
    try:
        from pptx import Presentation
    except Exception as e:
        return {"schema_version": 1, "error": f"python-pptx unavailable: {e}", "slides": []}

    if not pptx_bytes or not isinstance(pptx_bytes, (bytes, bytearray)):
        return {"schema_version": 1, "error": "empty_bytes", "slides": []}

    try:
        prs = Presentation(BytesIO(pptx_bytes))
    except Exception as e:
        return {"schema_version": 1, "error": str(e)[:200], "slides": []}

    sw = int(getattr(prs, "slide_width", 0) or 0)
    sh = int(getattr(prs, "slide_height", 0) or 0)
    if sw <= 0 or sh <= 0:
        return {
            "schema_version": 1,
            "slide_width_emu": sw,
            "slide_height_emu": sh,
            "slides": [],
            "error": "invalid_slide_dimensions",
        }

    slides_out: List[Dict[str, Any]] = []
    slide_errors: List[Dict[str, Any]] = []
    for si, slide in enumerate(prs.slides[:_MAX_SLIDES], start=1):
        try:
            shapes_out: List[Dict[str, Any]] = []
            for shape in list(slide.shapes)[:_MAX_SHAPES_PER_SLIDE]:
                try:
                    left = int(getattr(shape, "left", 0) or 0)
                    top = int(getattr(shape, "top", 0) or 0)
                    width = int(getattr(shape, "width", 0) or 0)
                    height = int(getattr(shape, "height", 0) or 0)
                except Exception:
                    continue
                bbox = [
                    round(left / sw, 4),
                    round(top / sh, 4),
                    round(width / sw, 4),
                    round(height / sh, 4),
                ]
                is_ph = False
                try:
                    is_ph = bool(getattr(shape, "is_placeholder", False))
                except Exception:
                    pass
                shapes_out.append(
                    {
                        "bbox": bbox,
                        "placeholder_type": _ph_type_name(shape),
                        "is_placeholder": is_ph,
                        "shape_kind": _shape_kind(shape),
                    }
                )
            slides_out.append({"index": si, "shapes": shapes_out})
        except Exception as e:
            slide_errors.append({"index": si, "error": str(e)[:300]})
            continue

    payload: Dict[str, Any] = {
        "schema_version": 1,
        "slide_width_emu": sw,
        "slide_height_emu": sh,
        "slide_count": len(prs.slides),
        "slides": slides_out,
        "slide_errors": slide_errors,
    }
    raw = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    if len(raw.encode("utf-8")) > _MAX_JSON_BYTES:
        payload["slides"] = slides_out[:24]
        payload["truncated"] = True
    return payload
