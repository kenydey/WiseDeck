from __future__ import annotations

from io import BytesIO

from pptx import Presentation
from pptx.util import Inches

from wisedeck.services.slide.pptx_roundtrip_bridge import pptx_bytes_to_pptist_slides


def test_pptx_roundtrip_bridge_smoke():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)  # 16:9

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "Hello"

    out = BytesIO()
    prs.save(out)
    out.seek(0)

    slides, meta = pptx_bytes_to_pptist_slides(out.read(), fixed_viewport=True)
    assert isinstance(slides, list) and slides
    assert meta["viewportSize"] == 1000.0
    assert slides[0]["background"]["type"] in ("solid", "image", "gradient")
    assert isinstance(slides[0]["elements"], list)

